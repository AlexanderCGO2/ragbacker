import os
from llama_parse import LlamaParse
from pydantic import BaseModel, validator
from pptx import Presentation
from llama_index.core.readers import SimpleDirectoryReader
import logging
# logging.basicConfig(level=logging.DEBUG)

class FileLoaderConfig(BaseModel):
    data_dir: str = "./data"
    use_llama_parse: bool = True

    @validator("data_dir")
    def data_dir_must_exist(cls, v):
        if not os.path.isdir(v):
            raise ValueError(f"Directory '{v}' does not exist")
        return v

def llama_parse_parser():
    if os.getenv("LLAMA_CLOUD_API_KEY") is None:
        raise ValueError(
            "LLAMA_CLOUD_API_KEY environment variable is not set. "
            "Please set it in .env file or in your shell environment then run again!"
        )
    parser = LlamaParse(result_type="markdown", verbose=True, language="de")
    return parser

# Erstellen Sie den Parser einmal und verwenden Sie ihn in der gesamten Anwendung
llama_parser = llama_parse_parser() if os.getenv("LLAMA_CLOUD_API_KEY") else None


def get_file_documents(config: FileLoaderConfig):
    reader = SimpleDirectoryReader(
        config.data_dir,
        recursive=True,
    )
    if config.use_llama_parse and llama_parser is not None:
        reader.file_extractor = {".pdf": llama_parser}
    else:
        # Unterstützung für verschiedene Dateitypen hinzufügen
        reader.file_extractor = {
            ".pptx": lambda path: {"text": [slide.text for slide in Presentation(path).slides], "filename": os.path.basename(path), "filetype": ".pptx"},
            ".md": lambda path: {"text": open(path).read(), "filename": os.path.basename(path), "filetype": ".md"},
            ".jpg": lambda path: {"text": path, "filename": os.path.basename(path), "filetype": ".jpg"},  # Beispiel: Pfad zurückgeben
            ".png": lambda path: {"text": path, "filename": os.path.basename(path), "filetype": ".png"},  # Beispiel: Pfad zurückgeben
            ".docx": lambda path: {"text": Document(path).text, "filename": os.path.basename(path), "filetype": ".docx"},  # Sie müssen die geeignete Bibliothek für die Verarbeitung von .docx-Dateien verwenden
            # Fügen Sie hier weitere Dateitypen hinzu...
        }
    return reader.load_data()